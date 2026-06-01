"""
Генераторы документов Word (.docx).

Поддерживаемые типы документов:
  - proposal  — Коммерческое предложение
  - contract  — Договор
  - act       — Акт выполненных работ
  - invoice   — Счёт на оплату
"""

import io
import os
from decimal import Decimal

from django.conf import settings
from django.utils import timezone

from docx import Document as DocxDocument
from docx.shared import Pt, Cm, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn


# ── Утилиты ─────────────────────────────────────────────────────────────────

def _fmt_money(val):
    if val is None:
        return '0.00 руб.'
    return f'{val:,.2f} руб.'.replace(',', ' ')


def _company_header():
    return {
        'name': 'ООО «БИТ»',
        'full': 'Общество с ограниченной ответственностью «Будущие Информационные Технологии»',
        'inn': '1644100856',
        'kpp': '164401001',
        'ogrn': '1221600080668',
        'address': 'респ. Татарстан (Татарстан), м.р-н Бугульминский, г. Бугульма, ул. Газинура Гафиатуллина, д. 16, кв. 44',
        'phone': '+7 (999) 123-45-67',
        'email': 'info@bit-company.ru',
        # Можно переопределить в settings.py: BIT_COMPANY_BANK, BIT_COMPANY_BIK,
        # BIT_COMPANY_CORR_ACCOUNT, BIT_COMPANY_SETTLEMENT_ACCOUNT
        'bank': getattr(settings, 'BIT_COMPANY_BANK', 'АО «ТБанк»'),
        'bik': getattr(settings, 'BIT_COMPANY_BIK', '044525974'),
        'corr_account': getattr(settings, 'BIT_COMPANY_CORR_ACCOUNT', '30101810145250000974'),
        'settlement_account': getattr(settings, 'BIT_COMPANY_SETTLEMENT_ACCOUNT', '40702810900000000001'),
        'director': 'Шакирзянов И.М.',
        'director_full': 'Шакирзянов Ильдар Мустакимович',
        'director_full_genitive': 'Шакирзянова Ильдара Мустакимовича',
        'director_position': 'Генеральный директор',
    }


def _get_client_company(biz_doc):
    """Получить данные компании клиента из заявки."""
    sr = biz_doc.request
    cc = getattr(sr, 'client_company', None)
    if cc:
        return {
            'name': cc.short_name,
            'full_name': cc.full_name or cc.short_name,
            'inn': cc.inn,
            'kpp': cc.kpp,
            'ogrn': cc.ogrn,
            'address': cc.legal_address,
            'bank': cc.bank_details_str,
            'director': cc.director_name,
            'director_genitive': cc.director_name,
            'director_position': cc.director_position,
            'director_basis': cc.director_basis,
            'phone': cc.phone or sr.contact_phone,
            'email': cc.email or sr.contact_email,
            'stamp_path': cc.stamp_image.path if cc.stamp_image else None,
            'signature_path': cc.signature_image.path if cc.signature_image else None,
        }
    # Fallback: собираем из заявки
    client = sr.client
    return {
        'name': sr.company_name or client.get_full_name() or client.username,
        'full_name': sr.company_name or client.get_full_name() or client.username,
        'inn': '',
        'kpp': '',
        'ogrn': '',
        'address': '',
        'bank': '',
        'director': client.get_full_name() or client.username,
        'director_genitive': client.get_full_name() or client.username,
        'director_position': '',
        'director_basis': '',
        'phone': sr.contact_phone,
        'email': sr.contact_email,
        'stamp_path': None,
        'signature_path': None,
    }


def _get_executor_stamp_path():
    """Путь к печати и подписи исполнителя (ООО БИТ)."""
    stamp = os.path.join(settings.MEDIA_ROOT, 'company', 'stamp.png')
    sig = os.path.join(settings.MEDIA_ROOT, 'company', 'signature.png')
    return (stamp if os.path.isfile(stamp) else None,
            sig if os.path.isfile(sig) else None)


# ═══════════════════════════════════════════════════════════════════════════
#  WORD (.docx) генераторы
# ═══════════════════════════════════════════════════════════════════════════

def _docx_set_style(doc):
    """Устанавливает базовый стиль документа."""
    style = doc.styles['Normal']
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(4)


def _docx_add_header(doc, company):
    """Добавляет шапку компании-исполнителя."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(company['name'])
    run.bold = True
    run.font.size = Pt(16)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run2 = p2.add_run(company['full'])
    run2.font.size = Pt(9)
    run2.font.color.rgb = RGBColor(100, 100, 100)

    p3 = doc.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run3 = p3.add_run(
        f'ИНН {company["inn"]} | КПП {company["kpp"]} | ОГРН {company["ogrn"]}\n'
        f'{company["address"]} | {company["phone"]} | {company["email"]}'
    )
    run3.font.size = Pt(8)
    run3.font.color.rgb = RGBColor(120, 120, 120)

    doc.add_paragraph('_' * 78)


def _docx_add_client_info(doc, biz_doc):
    """Добавляет полный блок с информацией о клиенте (включая реквизиты)."""
    client_data = _get_client_company(biz_doc)

    p = doc.add_paragraph()
    p.add_run('Заказчик: ').bold = True
    p.add_run(client_data['name'])

    if client_data['inn']:
        p2 = doc.add_paragraph()
        p2.add_run('ИНН: ').bold = True
        inn_str = client_data['inn']
        if client_data['kpp']:
            inn_str += f' / КПП: {client_data["kpp"]}'
        p2.add_run(inn_str)

    if client_data['address']:
        p3 = doc.add_paragraph()
        p3.add_run('Адрес: ').bold = True
        p3.add_run(client_data['address'])

    p4 = doc.add_paragraph()
    p4.add_run('Email: ').bold = True
    p4.add_run(client_data['email'])

    if client_data['phone']:
        p5 = doc.add_paragraph()
        p5.add_run('Телефон: ').bold = True
        p5.add_run(client_data['phone'])


def _docx_add_stamp_signature(cell_or_doc, image_path, width_inches=1.3):
    """Добавляет изображение печати/подписи."""
    if image_path and os.path.isfile(image_path):
        try:
            if hasattr(cell_or_doc, 'add_paragraph'):
                p = cell_or_doc.add_paragraph()
            else:
                p = cell_or_doc.paragraphs[0] if cell_or_doc.paragraphs else cell_or_doc.add_paragraph()
            run = p.add_run()
            run.add_picture(image_path, width=Inches(width_inches))
        except Exception:
            pass


def _docx_sign_block(doc, company, client_data, is_signed=False):
    """Добавляет таблицу с реквизитами сторон и отметкой об электронной подписи."""
    doc.add_paragraph()
    h = doc.add_paragraph()
    h.add_run('РЕКВИЗИТЫ СТОРОН').bold = True
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER

    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'

    # ── Исполнитель ──
    exec_lines = [
        'ИСПОЛНИТЕЛЬ:',
        company['name'],
        f'ИНН {company["inn"]} / КПП {company["kpp"]}',
        f'ОГРН {company["ogrn"]}',
        company['address'],
    ]
    if company.get('settlement_account'):
        exec_lines += [
            '',
            'Банковские реквизиты:',
            f'Р/с {company["settlement_account"]}',
            f'{company["bank"]}',
            f'БИК {company["bik"]}',
            f'К/с {company["corr_account"]}',
        ]
    if company.get('phone') or company.get('email'):
        exec_lines += ['', f'Контакты: {company.get("phone", "")} {company.get("email", "")}']
    if is_signed:
        exec_lines.append('Статус: подписано электронной подписью')

    cell_exec = table.rows[0].cells[0]
    cell_exec.text = '\n'.join(exec_lines)
    for paragraph in cell_exec.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(9)

    # ── Заказчик ──
    client_lines = ['ЗАКАЗЧИК:', client_data['name']]
    if client_data['inn']:
        inn_line = f'ИНН {client_data["inn"]}'
        if client_data['kpp']:
            inn_line += f' / КПП {client_data["kpp"]}'
        client_lines.append(inn_line)
    if client_data['ogrn']:
        client_lines.append(f'ОГРН {client_data["ogrn"]}')
    if client_data['address']:
        client_lines.append(client_data['address'])
    if client_data['bank']:
        client_lines += ['', 'Банковские реквизиты:', client_data['bank']]
    if client_data.get('phone') or client_data.get('email'):
        client_lines += ['', f'Контакты: {client_data.get("phone", "")} {client_data.get("email", "")}']
    if is_signed:
        client_lines.append('Статус: подписано электронной подписью')

    cell_client = table.rows[0].cells[1]
    cell_client.text = '\n'.join(client_lines)
    for paragraph in cell_client.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(9)


# ── Коммерческое предложение (Word) ─────────────────────────────────────

def generate_proposal_docx(biz_doc):
    """Генерирует КП в формате .docx и возвращает BytesIO."""
    doc = DocxDocument()
    _docx_set_style(doc)
    company = _company_header()
    client_data = _get_client_company(biz_doc)
    is_signed = biz_doc.status == 'signed'
    _docx_add_header(doc, company)

    # Заголовок
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(f'КОММЕРЧЕСКОЕ ПРЕДЛОЖЕНИЕ № {biz_doc.number}')
    run.bold = True
    run.font.size = Pt(14)

    doc.add_paragraph(f'от {(biz_doc.issue_date or timezone.now().date()).strftime("%d.%m.%Y")}')
    doc.add_paragraph()

    _docx_add_client_info(doc, biz_doc)
    doc.add_paragraph()

    sr = biz_doc.request
    doc.add_paragraph().add_run('Предмет предложения:').bold = True
    doc.add_paragraph(sr.title)
    doc.add_paragraph(sr.description)

    # Связанное КП с позициями
    if biz_doc.proposal and biz_doc.proposal.items.exists():
        doc.add_paragraph()
        doc.add_paragraph().add_run('Состав работ и стоимость:').bold = True
        items = biz_doc.proposal.items.all().order_by('order')
        table = doc.add_table(rows=1, cols=5)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = table.rows[0].cells
        for i, text in enumerate(['№', 'Этап', 'Кол-во', 'Цена', 'Сумма']):
            hdr[i].text = text
            hdr[i].paragraphs[0].runs[0].bold = True
        for idx, item in enumerate(items, 1):
            row = table.add_row().cells
            row[0].text = str(idx)
            row[1].text = item.stage_name
            row[2].text = str(item.quantity)
            row[3].text = _fmt_money(item.unit_price)
            row[4].text = _fmt_money(item.total_price)

    doc.add_paragraph()
    total_p = doc.add_paragraph()
    total_p.add_run(f'ИТОГО: {_fmt_money(biz_doc.amount)}').bold = True
    total_p.runs[0].font.size = Pt(13)

    if biz_doc.proposal and biz_doc.proposal.duration_days:
        doc.add_paragraph(f'Срок реализации: {biz_doc.proposal.duration_days} дней')

    if biz_doc.due_date:
        doc.add_paragraph(f'Предложение действительно до: {biz_doc.due_date.strftime("%d.%m.%Y")}')

    # Реквизиты сторон размещаем внизу документа
    _docx_sign_block(doc, company, client_data, is_signed=is_signed)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ── Договор (Word) ──────────────────────────────────────────────────────

def generate_contract_docx(biz_doc):
    """Генерирует договор в формате .docx."""
    doc = DocxDocument()
    _docx_set_style(doc)
    company = _company_header()
    client_data = _get_client_company(biz_doc)
    is_signed = biz_doc.status == 'signed'
    _docx_add_header(doc, company)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(f'ДОГОВОР № {biz_doc.number}')
    run.bold = True
    run.font.size = Pt(14)

    sr = biz_doc.request
    contract = biz_doc.contract
    date_str = (biz_doc.issue_date or timezone.now().date()).strftime('%d.%m.%Y')

    doc.add_paragraph(f'г. Бугульма                              «{date_str}»')
    doc.add_paragraph()

    # Стороны — с корректной юридической формулировкой
    executor_director_gen = company.get('director_full_genitive') or company.get('director_full')
    executor_basis = f'в лице генерального директора {executor_director_gen}, действующего на основании Устава'
    client_basis = ''
    client_director_gen = client_data.get('director_genitive') or client_data.get('director')
    if client_data['director'] and client_data['director_basis']:
        client_basis = (
            f'в лице генерального директора {client_director_gen}, '
            f'действующего на основании {client_data["director_basis"]}'
        )
    elif client_data['director']:
        client_basis = f'в лице генерального директора {client_director_gen}'

    doc.add_paragraph(
        f'{company["full"]}, именуемое в дальнейшем «Исполнитель», '
        f'{executor_basis}, с одной стороны, и '
        f'{client_data["full_name"]}, '
        f'{client_basis + ", " if client_basis else ""}'
        f'именуемый в дальнейшем «Заказчик», с другой стороны, '
        f'заключили настоящий Договор о нижеследующем:'
    )
    doc.add_paragraph()

    # 1. Предмет
    h1 = doc.add_paragraph()
    h1.add_run('1. ПРЕДМЕТ').bold = True
    subj = contract.subject if contract else sr.title
    doc.add_paragraph(f'1.1. Исполнитель обязуется оказать Заказчику услуги: {subj}.')
    doc.add_paragraph(f'1.2. Описание работ: {sr.description[:500]}')

    # 2. Стоимость
    doc.add_paragraph()
    h2 = doc.add_paragraph()
    h2.add_run('2. СТОИМОСТЬ И ПОРЯДОК РАСЧЁТОВ').bold = True
    doc.add_paragraph(f'2.1. Общая стоимость работ составляет {_fmt_money(biz_doc.amount)}.')
    if contract and contract.payment_terms:
        doc.add_paragraph(f'2.2. Условия оплаты: {contract.payment_terms}')
    else:
        doc.add_paragraph('2.2. Оплата производится в два этапа: 50% — предоплата, '
                          '50% — по факту подписания акта выполненных работ.')

    # 3. Сроки
    doc.add_paragraph()
    h3 = doc.add_paragraph()
    h3.add_run('3. СРОКИ ВЫПОЛНЕНИЯ РАБОТ').bold = True
    if contract and contract.start_date and contract.end_date:
        doc.add_paragraph(
            f'3.1. Начало работ: {contract.start_date.strftime("%d.%m.%Y")}. '
            f'Окончание работ: {contract.end_date.strftime("%d.%m.%Y")}.'
        )
    else:
        doc.add_paragraph('3.1. Сроки выполнения работ определяются согласно Техническому заданию.')

    # 4. Порядок сдачи-приёмки
    doc.add_paragraph()
    h4 = doc.add_paragraph()
    h4.add_run('4. ПОРЯДОК СДАЧИ-ПРИЁМКИ РАБОТ').bold = True
    doc.add_paragraph('4.1. По завершении работ Исполнитель предоставляет Заказчику Акт выполненных работ.')
    doc.add_paragraph('4.2. Заказчик обязуется рассмотреть Акт в течение 5 рабочих дней.')
    doc.add_paragraph('4.3. При наличии замечаний Заказчик направляет мотивированный отказ.')

    # 5. Ответственность
    doc.add_paragraph()
    h5 = doc.add_paragraph()
    h5.add_run('5. ОТВЕТСТВЕННОСТЬ СТОРОН').bold = True
    doc.add_paragraph('5.1. За неисполнение или ненадлежащее исполнение обязательств стороны '
                      'несут ответственность в соответствии с законодательством РФ.')
    doc.add_paragraph('5.2. Исполнитель гарантирует качество и соответствие работ ТЗ.')

    # 6. Конфиденциальность
    doc.add_paragraph()
    h6 = doc.add_paragraph()
    h6.add_run('6. КОНФИДЕНЦИАЛЬНОСТЬ').bold = True
    doc.add_paragraph('6.1. Стороны обязуются не разглашать конфиденциальную информацию, '
                      'полученную в ходе исполнения настоящего Договора.')

    # 7. Заключительные положения
    doc.add_paragraph()
    h7 = doc.add_paragraph()
    h7.add_run('7. ЗАКЛЮЧИТЕЛЬНЫЕ ПОЛОЖЕНИЯ').bold = True
    doc.add_paragraph('7.1. Настоящий Договор вступает в силу с момента подписания.')
    doc.add_paragraph('7.2. Составлен в двух экземплярах, имеющих одинаковую юридическую силу.')

    # Реквизиты сторон размещаем внизу документа
    _docx_sign_block(doc, company, client_data, is_signed=is_signed)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ── Акт выполненных работ (Word) ─────────────────────────────────────────

def generate_act_docx(biz_doc):
    """Генерирует акт выполненных работ в формате .docx."""
    doc = DocxDocument()
    _docx_set_style(doc)
    company = _company_header()
    client_data = _get_client_company(biz_doc)
    is_signed = biz_doc.status == 'signed'
    _docx_add_header(doc, company)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(f'АКТ ВЫПОЛНЕННЫХ РАБОТ № {biz_doc.number}')
    run.bold = True
    run.font.size = Pt(14)

    sr = biz_doc.request
    date_str = (biz_doc.issue_date or timezone.now().date()).strftime('%d.%m.%Y')
    doc.add_paragraph(f'от {date_str}')
    doc.add_paragraph()

    _docx_add_client_info(doc, biz_doc)
    doc.add_paragraph()

    contract_ref = ''
    if biz_doc.contract:
        contract_ref = f' по договору {biz_doc.contract.number}'

    doc.add_paragraph(
        f'Мы, нижеподписавшиеся, Исполнитель — {company["name"]}, и Заказчик — '
        f'{client_data["name"]}, составили настоящий Акт{contract_ref} '
        f'о том, что следующие работы выполнены в полном объёме и в установленные сроки:'
    )
    doc.add_paragraph()

    # Таблица работ
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    for i, text in enumerate(['№', 'Наименование работ', 'Объём', 'Сумма']):
        hdr[i].text = text
        hdr[i].paragraphs[0].runs[0].bold = True

    if biz_doc.proposal and biz_doc.proposal.items.exists():
        for idx, item in enumerate(biz_doc.proposal.items.all().order_by('order'), 1):
            row = table.add_row().cells
            row[0].text = str(idx)
            row[1].text = item.stage_name
            row[2].text = f'{item.quantity} {item.unit}'
            row[3].text = _fmt_money(item.total_price)
    else:
        row = table.add_row().cells
        row[0].text = '1'
        row[1].text = sr.title
        row[2].text = 'комплект'
        row[3].text = _fmt_money(biz_doc.amount)

    doc.add_paragraph()
    total_p = doc.add_paragraph()
    total_p.add_run(f'ИТОГО: {_fmt_money(biz_doc.amount)}').bold = True

    doc.add_paragraph()
    doc.add_paragraph('Вышеперечисленные работы выполнены полностью и в срок. '
                      'Заказчик претензий по объёму, качеству и срокам оказания услуг не имеет.')

    # Реквизиты сторон размещаем внизу документа
    _docx_sign_block(doc, company, client_data, is_signed=is_signed)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ── Счёт на оплату (Word) ────────────────────────────────────────────────

def generate_invoice_docx(biz_doc):
    """Генерирует счёт на оплату в формате .docx."""
    doc = DocxDocument()
    _docx_set_style(doc)
    company = _company_header()
    client_data = _get_client_company(biz_doc)
    is_signed = biz_doc.status == 'signed'
    _docx_add_header(doc, company)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(f'СЧЁТ НА ОПЛАТУ № {biz_doc.number}')
    run.bold = True
    run.font.size = Pt(14)

    sr = biz_doc.request
    date_str = (biz_doc.issue_date or timezone.now().date()).strftime('%d.%m.%Y')
    doc.add_paragraph(f'от {date_str}')
    doc.add_paragraph()

    # Реквизиты исполнителя (подробно)
    doc.add_paragraph().add_run('Получатель:').bold = True
    doc.add_paragraph(f'{company["name"]}  |  ИНН {company["inn"]}  |  КПП {company["kpp"]}')
    if company.get('settlement_account'):
        doc.add_paragraph(f'Р/с {company["settlement_account"]}')
    if company.get('bank') or company.get('bik'):
        bank_line = company.get('bank', '')
        if company.get('bik'):
            bank_line = f'{bank_line}, БИК {company["bik"]}' if bank_line else f'БИК {company["bik"]}'
        doc.add_paragraph(bank_line)
    if company.get('corr_account'):
        doc.add_paragraph(f'К/с {company["corr_account"]}')
    doc.add_paragraph()

    # Плательщик (полные реквизиты)
    doc.add_paragraph().add_run('Плательщик:').bold = True
    doc.add_paragraph(client_data['name'])
    if client_data['inn']:
        inn_str = f'ИНН {client_data["inn"]}'
        if client_data['kpp']:
            inn_str += f' / КПП {client_data["kpp"]}'
        doc.add_paragraph(inn_str)
    if client_data['address']:
        doc.add_paragraph(client_data['address'])
    doc.add_paragraph(f'Email: {client_data["email"]}')
    doc.add_paragraph()

    # Таблица
    table = doc.add_table(rows=1, cols=4)
    table.style = 'Table Grid'
    hdr = table.rows[0].cells
    for i, text in enumerate(['№', 'Наименование', 'Кол-во', 'Сумма']):
        hdr[i].text = text
        hdr[i].paragraphs[0].runs[0].bold = True

    if biz_doc.proposal and biz_doc.proposal.items.exists():
        for idx, item in enumerate(biz_doc.proposal.items.all().order_by('order'), 1):
            row = table.add_row().cells
            row[0].text = str(idx)
            row[1].text = item.stage_name
            row[2].text = f'{item.quantity} {item.unit}'
            row[3].text = _fmt_money(item.total_price)
    else:
        row = table.add_row().cells
        row[0].text = '1'
        row[1].text = sr.title
        row[2].text = '1'
        row[3].text = _fmt_money(biz_doc.amount)

    doc.add_paragraph()
    total_p = doc.add_paragraph()
    total_p.add_run(f'ИТОГО К ОПЛАТЕ: {_fmt_money(biz_doc.amount)}').bold = True
    total_p.runs[0].font.size = Pt(13)

    if biz_doc.due_date:
        doc.add_paragraph(f'Срок оплаты: до {biz_doc.due_date.strftime("%d.%m.%Y")}')

    doc.add_paragraph()
    doc.add_paragraph('Счёт действителен к оплате в течение 5 банковских дней.')

    # Отметка об электронной подписи
    doc.add_paragraph()
    doc.add_paragraph('')
    if is_signed:
        doc.add_paragraph('Статус: подписано электронной подписью')

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ═══════════════════════════════════════════════════════════════════════════
#  Диспетчер
# ═══════════════════════════════════════════════════════════════════════════

GENERATORS = {
    'proposal': generate_proposal_docx,
    'contract': generate_contract_docx,
    'act':      generate_act_docx,
    'invoice':  generate_invoice_docx,
}


def generate_document(biz_doc, fmt='docx'):
    """
    Генерирует документ в формате Word (.docx).
    Возвращает (BytesIO, content_type, filename).
    """
    gen = GENERATORS.get(biz_doc.doc_type)
    if not gen:
        raise ValueError(f'Неизвестный тип документа: {biz_doc.doc_type}')

    buf = gen(biz_doc)
    ct = 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    filename = f'{biz_doc.number}.docx'.replace('/', '-')
    return buf, ct, filename
