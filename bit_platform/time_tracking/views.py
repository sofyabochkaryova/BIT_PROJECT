from decimal import Decimal, InvalidOperation
from io import BytesIO

from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.utils import timezone

from notifications.views import create_notification
from tasks.models import Task
from .models import TimeEntry, WorkReport


# ───────────────────────── Учёт времени ─────────────────────────

@login_required
def log_time(request, task_id):
    """Исполнитель логирует время по задаче."""
    task = get_object_or_404(Task, pk=task_id)
    user = request.user
    profile = getattr(user, 'profile', None)
    role = getattr(profile, 'role', None)

    can_log = (
        role == 'executor' and (
            task.assignee == user
            or task.assignments.filter(assignee=user).exists()
        )
    )
    if not can_log:
        messages.error(request, 'Вы не назначены на эту задачу')
        return redirect('tasks:task_detail', pk=task_id)

    if request.method == 'POST':
        try:
            hours = Decimal(request.POST.get('hours', '0'))
        except (InvalidOperation, ValueError):
            hours = Decimal('0')

        description = request.POST.get('description', '').strip()
        date_str = request.POST.get('date', '')

        if hours <= 0:
            messages.error(request, 'Укажите корректное количество часов')
            return redirect('time_tracking:log_time', task_id=task_id)
        if not description:
            messages.error(request, 'Опишите проделанную работу')
            return redirect('time_tracking:log_time', task_id=task_id)

        entry = TimeEntry.objects.create(
            task=task,
            user=user,
            hours=hours,
            description=description,
            date=date_str or timezone.now().date(),
        )

        messages.success(request, f'Записано {hours} ч. по задаче «{task.title}»')
        return redirect('tasks:task_detail', pk=task_id)

    entries = TimeEntry.objects.filter(task=task, user=user).order_by('-date')
    total_mine = sum(e.hours for e in entries)

    return render(request, 'time_tracking/log_time.html', {
        'task': task,
        'entries': entries,
        'total_mine': total_mine,
    })


@login_required
def task_timesheet(request, task_id):
    """Все записи времени по задаче (для аналитика)."""
    task = get_object_or_404(Task, pk=task_id)
    entries = task.time_entries.select_related('user').order_by('-date')
    total = sum(e.hours for e in entries)

    return render(request, 'time_tracking/task_timesheet.html', {
        'task': task,
        'entries': entries,
        'total': total,
    })


# ───────────────────────── Отчёты о работе ─────────────────────────

@login_required
def create_work_report(request, task_id):
    """Исполнитель создаёт отчёт о проделанной работе."""
    task = get_object_or_404(Task, pk=task_id)
    user = request.user
    profile = getattr(user, 'profile', None)
    role = getattr(profile, 'role', None)

    can_create_report = (
        role == 'executor' and (
            task.assignee == user
            or task.assignments.filter(assignee=user).exists()
        )
    )
    if not can_create_report:
        messages.error(request, 'Отчёт по задаче может создавать только назначенный исполнитель')
        return redirect('tasks:task_detail', pk=task_id)

    if request.method == 'POST':
        title = request.POST.get('title', '').strip()
        content = request.POST.get('content', '').strip()

        if not title or not content:
            messages.error(request, 'Заполните заголовок и содержание')
            return redirect('time_tracking:create_report', task_id=task_id)

        # Считаем суммарные часы
        from django.db.models import Sum
        hours_total = task.time_entries.filter(user=user).aggregate(s=Sum('hours'))['s'] or 0

        report = WorkReport.objects.create(
            task=task,
            author=user,
            title=title,
            content=content,
            hours_total=hours_total,
            status='submitted',
        )

        # Уведомляем аналитика / менеджера проекта
        manager = task.project.manager
        if manager and manager != user:
            create_notification(
                recipient=manager,
                title='Новый отчёт о работе',
                message=f'{user.get_full_name() or user.username} отправил отчёт по задаче «{task.title}»',
                notification_type='request',
                sender=user,
                link=f'/time-tracking/report/{report.pk}/',
            )

        messages.success(request, 'Отчёт отправлен на проверку')
        return redirect('time_tracking:report_detail', pk=report.pk)

    # Собираем записи времени для контекста
    entries = task.time_entries.filter(user=user).order_by('-date')
    total_hours = sum(e.hours for e in entries)

    return render(request, 'time_tracking/create_report.html', {
        'task': task,
        'entries': entries,
        'total_hours': total_hours,
    })


@login_required
def report_detail(request, pk):
    """Просмотр отчёта о работе."""
    report = get_object_or_404(
        WorkReport.objects.select_related('task', 'task__project', 'author', 'reviewer'),
        pk=pk,
    )

    user = request.user
    profile = getattr(user, 'profile', None)
    role = getattr(profile, 'role', None)

    is_reviewer = role in ['analyst', 'admin'] or report.task.project.manager == user
    is_author = report.author == user

    if not (is_author or is_reviewer):
        messages.error(request, 'Нет доступа к отчёту')
        return redirect('accounts:dashboard')

    if request.method == 'POST' and is_reviewer:
        action = request.POST.get('action')

        if action == 'accept':
            report.status = 'accepted'
            report.reviewer = user
            report.review_comment = request.POST.get('comment', '')
            report.reviewed_at = timezone.now()
            report.save()
            create_notification(
                recipient=report.author,
                title='Отчёт принят',
                message=f'Ваш отчёт «{report.title}» принят',
                notification_type='request',
                sender=user,
                link=f'/time-tracking/report/{report.pk}/',
            )
            messages.success(request, 'Отчёт принят')

        elif action == 'revision':
            report.status = 'revision'
            report.reviewer = user
            report.review_comment = request.POST.get('comment', '')
            report.reviewed_at = timezone.now()
            report.save()
            create_notification(
                recipient=report.author,
                title='Отчёт на доработке',
                message=f'Отчёт «{report.title}» отправлен на доработку: {report.review_comment[:100]}',
                notification_type='request',
                sender=user,
                link=f'/time-tracking/report/{report.pk}/',
            )
            messages.success(request, 'Отчёт отправлен на доработку')

        return redirect('time_tracking:report_detail', pk=pk)

    # Если автор пере-отправляет после доработки
    if request.method == 'POST' and is_author and report.status == 'revision':
        content = request.POST.get('content', '').strip()
        if content:
            report.content = content
            report.status = 'submitted'
            report.save()
            if report.reviewer:
                create_notification(
                    recipient=report.reviewer,
                    title='Отчёт исправлен',
                    message=f'{report.author.get_full_name()} исправил отчёт «{report.title}»',
                    notification_type='request',
                    sender=user,
                    link=f'/time-tracking/report/{report.pk}/',
                )
            messages.success(request, 'Отчёт отправлен повторно')
        return redirect('time_tracking:report_detail', pk=pk)

    # Записи времени к задаче
    entries = report.task.time_entries.filter(user=report.author).order_by('-date')

    return render(request, 'time_tracking/report_detail.html', {
        'report': report,
        'is_reviewer': is_reviewer,
        'is_author': is_author,
        'entries': entries,
    })


@login_required
def my_reports(request):
    """Список отчётов текущего пользователя (исполнитель) или на проверке (аналитик)."""
    user = request.user
    profile = getattr(user, 'profile', None)
    role = getattr(profile, 'role', None)

    if role in ['analyst', 'admin']:
        reports = WorkReport.objects.select_related('task', 'author').order_by('-created_at')
    else:
        reports = WorkReport.objects.filter(author=user).select_related('task').order_by('-created_at')

    status_filter = request.GET.get('status')
    if status_filter:
        reports = reports.filter(status=status_filter)

    return render(request, 'time_tracking/my_reports.html', {
        'reports': reports,
        'statuses': WorkReport.STATUS_CHOICES,
        'current_status': status_filter,
    })


@login_required
def download_report_docx(request, pk):
    """Скачать отчёт о работе в формате Word."""
    report = get_object_or_404(
        WorkReport.objects.select_related('task', 'task__project', 'author'),
        pk=pk,
    )

    try:
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        messages.error(request, 'Модуль python-docx не установлен')
        return redirect('time_tracking:report_detail', pk=pk)

    doc = Document()
    style = doc.styles['Normal']
    style.font.size = Pt(11)
    style.font.name = 'Calibri'

    doc.add_heading('Отчёт о выполненных работах', level=1)
    doc.add_paragraph(f'Задача: {report.task.title}')
    doc.add_paragraph(f'Проект: {report.task.project.name}')
    doc.add_paragraph(f'Исполнитель: {report.author.get_full_name() or report.author.username}')
    doc.add_paragraph(f'Дата: {report.created_at.strftime("%d.%m.%Y")}')
    doc.add_paragraph(f'Затрачено часов: {report.hours_total}')
    doc.add_paragraph('')
    doc.add_heading(report.title, level=2)
    doc.add_paragraph(report.content)

    # Таблица записей времени
    entries = report.task.time_entries.filter(user=report.author).order_by('date')
    if entries.exists():
        doc.add_heading('Детализация времени', level=2)
        table = doc.add_table(rows=1, cols=3, style='Table Grid')
        table.rows[0].cells[0].text = 'Дата'
        table.rows[0].cells[1].text = 'Часы'
        table.rows[0].cells[2].text = 'Описание'
        for e in entries:
            row = table.add_row()
            row.cells[0].text = e.date.strftime('%d.%m.%Y')
            row.cells[1].text = str(e.hours)
            row.cells[2].text = e.description

    if report.review_comment:
        doc.add_heading('Комментарий проверяющего', level=2)
        doc.add_paragraph(report.review_comment)

    buf = BytesIO()
    doc.save(buf)
    buf.seek(0)

    response = HttpResponse(
        buf.read(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',
    )
    response['Content-Disposition'] = f'attachment; filename="report_{report.pk}.docx"'
    return response
