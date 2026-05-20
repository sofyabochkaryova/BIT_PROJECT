import io
from datetime import timedelta

from django.core.files.base import ContentFile
from django.shortcuts import render, get_object_or_404, redirect
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse
from django.db.models import Count, Q, Avg, Sum
from django.db.models.functions import TruncDate, TruncMonth
from django.utils import timezone
from django.utils.text import slugify

from requests_app.models import ServiceRequest
from projects.models import Project
from tasks.models import Task
from django.contrib.auth import get_user_model
from .models import Report, ActivityLog

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

User = get_user_model()


@login_required
def analytics_dashboard(request):
    """Главный дашборд аналитики"""
    user = request.user
    today = timezone.now().date()
    month_start = today.replace(day=1)
    week_start = today - timedelta(days=today.weekday())
    
    # Базовые фильтры в зависимости от роли
    if hasattr(user, 'profile') and user.profile.is_client:
        requests_qs = ServiceRequest.objects.filter(client=user)
        projects_qs = Project.objects.filter(client=user)
        tasks_qs = Task.objects.none()
    elif hasattr(user, 'profile') and user.profile.is_analyst:
        requests_qs = ServiceRequest.objects.filter(analyst=user)
        projects_qs = Project.objects.filter(manager=user)
        tasks_qs = Task.objects.filter(Q(created_by=user) | Q(project__manager=user))
    elif hasattr(user, 'profile') and user.profile.is_executor:
        requests_qs = ServiceRequest.objects.none()
        projects_qs = Project.objects.filter(team=user)
        tasks_qs = Task.objects.filter(assignee=user)
    else:  # Admin
        requests_qs = ServiceRequest.objects.all()
        projects_qs = Project.objects.all()
        tasks_qs = Task.objects.all()
    
    # Статистика
    stats = {
        'total_requests': requests_qs.count(),
        'new_requests': requests_qs.filter(status__in=['IN_ANALYSIS', 'ESTIMATION', 'PROPOSAL_DRAFT']).count(),
        'in_progress_requests': requests_qs.filter(status__in=['IN_PROGRESS', 'CONVERTED']).count(),
        'completed_requests': requests_qs.filter(status='COMPLETED').count(),
        
        'total_projects': projects_qs.count(),
        'active_projects': projects_qs.filter(status='active').count(),
        
        'total_tasks': tasks_qs.count(),
        'pending_tasks': tasks_qs.filter(status__in=['backlog', 'todo']).count(),
        'in_progress_tasks': tasks_qs.filter(status='in_progress').count(),
        'completed_tasks': tasks_qs.filter(status='done').count(),
        
        # За месяц
        'requests_this_month': requests_qs.filter(created_at__gte=month_start).count(),
        'completed_this_month': requests_qs.filter(
            completed_at__gte=month_start, status='COMPLETED'
        ).count(),
        
        # За неделю
        'requests_this_week': requests_qs.filter(created_at__gte=week_start).count(),
        'tasks_this_week': tasks_qs.filter(created_at__gte=week_start).count(),
    }
    
    # Графики - заявки по дням за последние 30 дней
    thirty_days_ago = today - timedelta(days=30)
    requests_by_day = list(
        requests_qs.filter(created_at__date__gte=thirty_days_ago)
        .annotate(date=TruncDate('created_at'))
        .values('date')
        .annotate(count=Count('id'))
        .order_by('date')
    )
    
    # Заявки по статусам
    requests_by_status = list(
        requests_qs.values('status')
        .annotate(count=Count('id'))
    )
    
    # Задачи по приоритетам
    tasks_by_priority = list(
        tasks_qs.values('priority')
        .annotate(count=Count('id'))
    )
    
    # Последняя активность
    recent_activity = ActivityLog.objects.filter(user=user)[:10]
    
    context = {
        'stats': stats,
        'requests_by_day': requests_by_day,
        'requests_by_status': requests_by_status,
        'tasks_by_priority': tasks_by_priority,
        'recent_activity': recent_activity,
    }
    
    return render(request, 'reports/analytics_dashboard.html', context)


@login_required
def reports_list(request):
    """Список отчётов"""
    reports = Report.objects.filter(created_by=request.user)
    
    return render(request, 'reports/reports_list.html', {'reports': reports})


@login_required
def create_report(request):
    """Создать новый отчёт"""
    if request.method == 'POST':
        date_from = request.POST.get('date_from') or None
        date_to = request.POST.get('date_to') or None

        report = Report.objects.create(
            name=request.POST.get('name'),
            report_type=request.POST.get('report_type'),
            description=request.POST.get('description', ''),
            date_from=date_from,
            date_to=date_to,
            created_by=request.user
        )
        report.refresh_from_db()  # Гарантирует, что date_from/date_to — объекты date

        headers, rows = _build_report_rows(report)
        excel_bytes = _build_excel_file(headers, rows, report)
        file_name = _generate_report_filename(report)
        report.file_format = 'xlsx'
        report.file.save(file_name, ContentFile(excel_bytes), save=True)

        log_activity(
            user=request.user,
            action='create',
            content_type='report',
            object_id=report.pk,
            object_repr=report.name,
            details={'report_type': report.report_type, 'rows': len(rows)},
            request=request,
        )
        
        if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
            return JsonResponse({
                'success': True,
                'report_id': report.pk
            })
        
        return redirect('reports:detail', pk=report.pk)
    
    return render(request, 'reports/create_report.html')


@login_required
def report_detail(request, pk):
    """Детали отчёта"""
    report = get_object_or_404(Report, pk=pk, created_by=request.user)

    preview_headers, preview_rows = _build_report_rows(report, limit=50)
    return render(
        request,
        'reports/report_detail.html',
        {
            'report': report,
            'preview_headers': preview_headers,
            'preview_rows': preview_rows,
        },
    )


@login_required
def delete_report(request, pk):
    """Удалить отчёт"""
    report = get_object_or_404(Report, pk=pk, created_by=request.user)

    if report.file:
        report.file.delete(save=False)

    log_activity(
        user=request.user,
        action='delete',
        content_type='report',
        object_id=report.pk,
        object_repr=report.name,
        request=request,
    )
    report.delete()
    
    if request.headers.get('X-Requested-With') == 'XMLHttpRequest':
        return JsonResponse({'success': True})
    
    return redirect('reports:list')


# API endpoints для графиков
@login_required
def api_requests_stats(request):
    """API: Статистика по заявкам"""
    period = request.GET.get('period', '30')  # дней
    
    user = request.user
    if hasattr(user, 'profile') and user.profile.is_client:
        qs = ServiceRequest.objects.filter(client=user)
    elif hasattr(user, 'profile') and user.profile.is_analyst:
        qs = ServiceRequest.objects.filter(analyst=user)
    else:
        qs = ServiceRequest.objects.all()
    
    start_date = timezone.now().date() - timedelta(days=int(period))
    
    # По дням
    by_day = list(
        qs.filter(created_at__date__gte=start_date)
        .annotate(date=TruncDate('created_at'))
        .values('date')
        .annotate(count=Count('id'))
        .order_by('date')
    )
    
    # По статусам
    by_status = list(
        qs.values('status')
        .annotate(count=Count('id'))
    )
    
    # По типам услуг
    by_service = list(
        qs.values('service_type')
        .annotate(count=Count('id'))
        .order_by('-count')[:5]
    )
    
    return JsonResponse({
        'by_day': [{'date': d['date'].isoformat(), 'count': d['count']} for d in by_day],
        'by_status': by_status,
        'by_service': by_service,
    })


@login_required
def api_tasks_stats(request):
    """API: Статистика по задачам"""
    user = request.user
    
    if hasattr(user, 'profile') and user.profile.is_executor:
        qs = Task.objects.filter(assignee=user)
    elif hasattr(user, 'profile') and user.profile.is_analyst:
        qs = Task.objects.filter(Q(created_by=user) | Q(project__manager=user))
    else:
        qs = Task.objects.all()
    
    # По статусам
    by_status = list(
        qs.values('status')
        .annotate(count=Count('id'))
    )
    
    # По приоритетам
    by_priority = list(
        qs.values('priority')
        .annotate(count=Count('id'))
    )
    
    # Просроченные
    overdue = qs.filter(
        due_date__lt=timezone.now().date(),
        status__in=['backlog', 'todo', 'in_progress']
    ).count()
    
    return JsonResponse({
        'by_status': by_status,
        'by_priority': by_priority,
        'overdue': overdue,
        'total': qs.count(),
    })


@login_required
def api_projects_stats(request):
    """API: Статистика по проектам"""
    user = request.user
    
    if hasattr(user, 'profile') and user.profile.is_client:
        qs = Project.objects.filter(client=user)
    elif hasattr(user, 'profile') and user.profile.is_executor:
        qs = Project.objects.filter(team=user)
    else:
        qs = Project.objects.all()
    
    # По статусам
    by_status = list(
        qs.values('status')
        .annotate(count=Count('id'))
    )
    
    # Средний прогресс
    avg_progress = qs.filter(status='active').aggregate(avg=Avg('progress'))
    
    return JsonResponse({
        'by_status': by_status,
        'avg_progress': avg_progress['avg'] or 0,
        'total': qs.count(),
    })


@login_required
def api_employees_stats(request):
    """API: Статистика по сотрудникам (эффективность)"""
    executors = User.objects.filter(profile__role='executor')
    data = []
    for emp in executors:
        total = Task.objects.filter(assignee=emp).count()
        done = Task.objects.filter(assignee=emp, status='done').count()
        overdue = Task.objects.filter(
            assignee=emp,
            due_date__lt=timezone.now().date(),
            status__in=['backlog', 'todo', 'in_progress']
        ).count()
        hours = 0
        try:
            from time_tracking.models import TimeEntry
            hours = float(
                TimeEntry.objects.filter(user=emp).aggregate(s=Sum('hours'))['s'] or 0
            )
        except Exception:
            pass
        data.append({
            'name': emp.get_full_name() or emp.username,
            'total_tasks': total,
            'done_tasks': done,
            'overdue_tasks': overdue,
            'efficiency': round(done / total * 100, 1) if total else 0,
            'hours_logged': round(hours, 1),
        })
    data.sort(key=lambda x: x['efficiency'], reverse=True)
    return JsonResponse({'employees': data})


@login_required
def api_clients_stats(request):
    """API: Статистика по клиентам"""
    clients = User.objects.filter(profile__role='client')
    data = []
    for c in clients:
        reqs = ServiceRequest.objects.filter(client=c)
        total = reqs.count()
        if total == 0:
            continue
        completed = reqs.filter(status='COMPLETED').count()
        in_progress = reqs.filter(status__in=['IN_PROGRESS', 'CONVERTED']).count()
        projects_count = Project.objects.filter(client=c).count()
        data.append({
            'name': c.get_full_name() or c.username,
            'total_requests': total,
            'completed': completed,
            'in_progress': in_progress,
            'projects': projects_count,
            'completion_rate': round(completed / total * 100, 1) if total else 0,
        })
    data.sort(key=lambda x: x['total_requests'], reverse=True)
    return JsonResponse({'clients': data})


# Утилита для логирования активности
def log_activity(user, action, content_type, object_id=None, object_repr='', details=None, request=None):
    """Записать активность пользователя"""
    ip_address = None
    user_agent = ''
    
    if request:
        x_forwarded_for = request.META.get('HTTP_X_FORWARDED_FOR')
        if x_forwarded_for:
            ip_address = x_forwarded_for.split(',')[0]
        else:
            ip_address = request.META.get('REMOTE_ADDR')
        user_agent = request.META.get('HTTP_USER_AGENT', '')[:500]
    
    ActivityLog.objects.create(
        user=user,
        action=action,
        content_type=content_type,
        object_id=object_id,
        object_repr=object_repr,
        details=details or {},
        ip_address=ip_address,
        user_agent=user_agent
    )


def _build_excel_file(headers, rows, report=None):
    """Генерация красиво оформленного Excel-файла с openpyxl."""
    wb = openpyxl.Workbook()
    ws = wb.active

    # Название типа отчёта
    type_names = {
        'requests': 'Заявки',
        'projects': 'Проекты',
        'tasks': 'Задачи',
        'users': 'Сотрудники',
        'financial': 'Клиенты',
    }
    ws.title = type_names.get(report.report_type, 'Отчёт') if report else 'Отчёт'

    # ── Стили ──
    header_font = Font(name='Calibri', bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill(start_color='2B579A', end_color='2B579A', fill_type='solid')
    header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    data_font = Font(name='Calibri', size=10)
    data_alignment = Alignment(vertical='center', wrap_text=True)

    stripe_fill = PatternFill(start_color='F2F6FC', end_color='F2F6FC', fill_type='solid')

    thin_border = Border(
        left=Side(style='thin', color='D0D5DD'),
        right=Side(style='thin', color='D0D5DD'),
        top=Side(style='thin', color='D0D5DD'),
        bottom=Side(style='thin', color='D0D5DD'),
    )

    # ── Заголовок отчёта (строка 1) ──
    title_text = report.name if report else 'Отчёт'
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
    title_cell = ws.cell(row=1, column=1, value=title_text)
    title_cell.font = Font(name='Calibri', bold=True, size=14, color='2B579A')
    title_cell.alignment = Alignment(horizontal='center', vertical='center')

    # ── Метаинформация ──
    meta_row = 2
    if report:
        period_text = ''
        if report.date_from and report.date_to:
            period_text = f'Период: {report.date_from.strftime("%d.%m.%Y")} — {report.date_to.strftime("%d.%m.%Y")}'
        elif report.date_from:
            period_text = f'С {report.date_from.strftime("%d.%m.%Y")}'
        elif report.date_to:
            period_text = f'По {report.date_to.strftime("%d.%m.%Y")}'
        else:
            period_text = 'За всё время'
        ws.merge_cells(start_row=meta_row, start_column=1, end_row=meta_row, end_column=len(headers))
        meta_cell = ws.cell(row=meta_row, column=1, value=period_text)
        meta_cell.font = Font(name='Calibri', size=10, italic=True, color='666666')
        meta_cell.alignment = Alignment(horizontal='center')
        meta_row += 1

    ws.merge_cells(start_row=meta_row, start_column=1, end_row=meta_row, end_column=len(headers))
    gen_cell = ws.cell(row=meta_row, column=1,
                       value=f'Сформировано: {timezone.now().strftime("%d.%m.%Y %H:%M")}  |  Записей: {len(rows)}')
    gen_cell.font = Font(name='Calibri', size=9, italic=True, color='999999')
    gen_cell.alignment = Alignment(horizontal='center')

    header_row_num = meta_row + 1  # пустая строка перед заголовками
    header_row_num += 1  # сами заголовки

    # ── Заголовки таблицы ──
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=header_row_num, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = thin_border

    # ── Данные ──
    for row_idx, row_data in enumerate(rows):
        excel_row = header_row_num + 1 + row_idx
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=excel_row, column=col_idx, value=value)
            cell.font = data_font
            cell.alignment = data_alignment
            cell.border = thin_border
            # Чередование строк
            if row_idx % 2 == 1:
                cell.fill = stripe_fill

    # ── Итоговая строка ──
    if rows:
        total_row = header_row_num + 1 + len(rows) + 1
        total_fill = PatternFill(start_color='E8EDF5', end_color='E8EDF5', fill_type='solid')
        total_font = Font(name='Calibri', bold=True, size=10)
        ws.cell(row=total_row, column=1, value=f'Итого записей: {len(rows)}')
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=total_row, column=col_idx)
            cell.fill = total_fill
            cell.font = total_font
            cell.border = thin_border

    # ── Автоширина столбцов ──
    for col_idx, header in enumerate(headers, 1):
        max_len = len(str(header))
        for row_data in rows[:100]:  # Сэмпл первых 100 строк
            if col_idx - 1 < len(row_data):
                val_len = len(str(row_data[col_idx - 1]))
                if val_len > max_len:
                    max_len = val_len
        adjusted_width = min(max_len + 4, 50)
        ws.column_dimensions[get_column_letter(col_idx)].width = adjusted_width

    # ── Закрепление заголовков ──
    ws.freeze_panes = ws.cell(row=header_row_num + 1, column=1)

    # ── Автофильтр ──
    last_col_letter = get_column_letter(len(headers))
    last_data_row = header_row_num + len(rows)
    ws.auto_filter.ref = f'A{header_row_num}:{last_col_letter}{last_data_row}'

    # ── Запись в буфер ──
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _generate_report_filename(report):
    safe_name = slugify(report.name) or f'report-{report.pk}'
    timestamp = timezone.now().strftime('%Y%m%d_%H%M%S')
    return f'{safe_name}_{timestamp}.xlsx'


def _scoped_querysets(user):
    if hasattr(user, 'profile') and user.profile.is_client:
        return {
            'requests': ServiceRequest.objects.filter(client=user),
            'projects': Project.objects.filter(client=user),
            'tasks': Task.objects.filter(assignee=user),
        }
    if hasattr(user, 'profile') and user.profile.is_analyst:
        return {
            'requests': ServiceRequest.objects.filter(Q(analyst=user) | Q(client=user)).distinct(),
            'projects': Project.objects.filter(Q(manager=user) | Q(client=user)).distinct(),
            'tasks': Task.objects.filter(Q(created_by=user) | Q(project__manager=user)).distinct(),
        }
    if hasattr(user, 'profile') and user.profile.is_executor:
        return {
            'requests': ServiceRequest.objects.filter(executor=user),
            'projects': Project.objects.filter(team=user),
            'tasks': Task.objects.filter(assignee=user),
        }
    return {
        'requests': ServiceRequest.objects.all(),
        'projects': Project.objects.all(),
        'tasks': Task.objects.all(),
    }


def _apply_date_filters(queryset, field_name, date_from, date_to):
    if date_from:
        queryset = queryset.filter(**{f'{field_name}__date__gte': date_from})
    if date_to:
        queryset = queryset.filter(**{f'{field_name}__date__lte': date_to})
    return queryset


def _build_report_rows(report, limit=None):
    scoped = _scoped_querysets(report.created_by)

    if report.report_type == 'requests':
        queryset = _apply_date_filters(scoped['requests'], 'created_at', report.date_from, report.date_to)
        queryset = queryset.select_related('client').order_by('-created_at')
        if limit:
            queryset = queryset[:limit]
        headers = [
            'ID', 'Название', 'Тип услуги', 'Статус', 'Приоритет', 'Клиент',
            'Компания', 'Дата создания',
        ]
        rows = [
            [
                item.pk,
                item.title,
                item.get_service_type_display(),
                item.get_status_display(),
                item.get_priority_display(),
                item.client.get_full_name() or item.client.username,
                item.company_name or '-',
                item.created_at.strftime('%d.%m.%Y %H:%M'),
            ]
            for item in queryset
        ]
        return headers, rows

    if report.report_type == 'projects':
        queryset = _apply_date_filters(scoped['projects'], 'created_at', report.date_from, report.date_to)
        queryset = queryset.select_related('client', 'manager').order_by('-created_at')
        if limit:
            queryset = queryset[:limit]
        headers = [
            'ID', 'Название', 'Статус', 'Прогресс %', 'Клиент', 'Менеджер',
            'Бюджет', 'Дата начала', 'Дата окончания', 'Дата создания',
        ]
        rows = [
            [
                item.pk,
                item.name,
                item.get_status_display(),
                item.progress,
                item.client.get_full_name() or item.client.username,
                (item.manager.get_full_name() or item.manager.username) if item.manager else '-',
                item.budget or '-',
                item.start_date.strftime('%d.%m.%Y') if item.start_date else '-',
                item.end_date.strftime('%d.%m.%Y') if item.end_date else '-',
                item.created_at.strftime('%d.%m.%Y %H:%M'),
            ]
            for item in queryset
        ]
        return headers, rows

    if report.report_type == 'users':
        # Отчёт по сотрудникам (эффективность)
        executors = User.objects.filter(profile__role='executor')
        headers = [
            'Сотрудник', 'Email', 'Всего задач', 'Выполнено', 'Просрочено',
            'Эффективность %', 'Активных проектов',
        ]
        rows = []
        for emp in executors:
            total_t = Task.objects.filter(assignee=emp).count()
            done_t = Task.objects.filter(assignee=emp, status='done').count()
            overdue_t = Task.objects.filter(
                assignee=emp, due_date__lt=timezone.now().date(),
                status__in=['backlog', 'todo', 'in_progress']
            ).count()
            hours = 0
            try:
                from time_tracking.models import TimeEntry
                hours = float(TimeEntry.objects.filter(user=emp).aggregate(s=Sum('hours'))['s'] or 0)
            except Exception:
                pass
            active_proj = Project.objects.filter(team=emp, status='active').count()
            rows.append([
                emp.get_full_name() or emp.username,
                emp.email,
                total_t, done_t, overdue_t,
                round(done_t / total_t * 100, 1) if total_t else 0,
                active_proj,
            ])
        return headers, rows

    if report.report_type == 'financial':
        # Отчёт по клиентам
        clients = User.objects.filter(profile__role='client')
        headers = [
            'Клиент', 'Email', 'Компания', 'Заявок всего', 'Завершено',
            'В работе', 'Проектов', '% завершения',
        ]
        rows = []
        for c in clients:
            reqs = ServiceRequest.objects.filter(client=c)
            total_r = reqs.count()
            if total_r == 0:
                continue
            completed_r = reqs.filter(status='COMPLETED').count()
            in_progress_r = reqs.filter(status__in=['IN_PROGRESS', 'CONVERTED']).count()
            projects_cnt = Project.objects.filter(client=c).count()
            company_name = reqs.first().company_name or '-'
            rows.append([
                c.get_full_name() or c.username,
                c.email,
                company_name,
                total_r, completed_r, in_progress_r,
                projects_cnt,
                round(completed_r / total_r * 100, 1) if total_r else 0,
            ])
        return headers, rows

    # default: tasks
    queryset = _apply_date_filters(scoped['tasks'], 'created_at', report.date_from, report.date_to)
    queryset = queryset.select_related('project', 'assignee', 'created_by').order_by('-created_at')
    if limit:
        queryset = queryset[:limit]
    headers = [
        'ID', 'Задача', 'Проект', 'Тип', 'Статус', 'Приоритет',
        'Исполнитель', 'Автор', 'Оценка (ч)', 'Затрачено (ч)', 'Дедлайн', 'Дата создания',
    ]
    rows = [
        [
            item.pk,
            item.title,
            item.project.name,
            item.get_task_type_display(),
            item.get_status_display(),
            item.get_priority_display(),
            (item.assignee.get_full_name() or item.assignee.username) if item.assignee else '-',
            item.created_by.get_full_name() or item.created_by.username,
            item.estimated_hours or '-',
            item.spent_hours or 0,
            item.due_date.strftime('%d.%m.%Y') if item.due_date else '-',
            item.created_at.strftime('%d.%m.%Y %H:%M'),
        ]
        for item in queryset
    ]
    return headers, rows
